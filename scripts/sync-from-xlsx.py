"""
sync-from-xlsx.py — regenerate src/lib/data/seed.ts from "Demo Inventory Control.xlsx".

What it syncs from Excel:
  - items[]   ← Component Inventory Tracker, rows 50+
  - ledgers   ← one per-item sheet per item ID (batched lot/location entries)

What it preserves from the existing seed.ts (not in Excel):
  - roles, users, auditEvents, signoffs

Usage:
  python scripts/sync-from-xlsx.py [--xlsx PATH]

Threshold policy: Excel-literal (col M, MUST HAVE INVENTORY LEVEL). Many items
will end up threshold=0, which is faithful to the source.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl not installed. Run: python -m pip install openpyxl\n")
    sys.exit(1)


REPO_ROOT = Path(__file__).resolve().parent.parent
SEED_PATH = REPO_ROOT / "frontend" / "src" / "lib" / "data" / "seed.ts"
DEFAULT_XLSX = Path(r"C:\Users\inventory\Desktop\Projects\inventory-claudeai\Demo Inventory Control.xlsx")

CATEGORY_LABEL = {
    "CM": "Cell-therapy media / reagent",
    "NM": "Manufacturing reagent",
    "CR": "Consumable / equipment",
    "AY": "Assay / QC reference",
}

MASTER_HEADER_ROW = 49  # 1-based
ID_RE = re.compile(r"^(CM|NM|CR|AY)-\d{4}$")


def iso_date(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s or s in ("—", "-", "Not in MC"):
        return None
    return s


def to_int(v) -> int:
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (v != v):  # NaN
            return 0
        return int(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    if not s or s in ("—", "-"):
        return 0
    try:
        return int(float(s))
    except ValueError:
        return 0


def derive_status(qty: int, threshold: int, expires_iso: str | None, today: date) -> str:
    if expires_iso:
        try:
            exp = datetime.strptime(expires_iso, "%Y-%m-%d").date()
            days = (exp - today).days
            if days < 0:
                return "Expired"
            if days <= 90 and qty > 0:
                return "Expiring"
        except ValueError:
            pass
    if qty <= 0:
        return "Out of stock"
    if threshold > 0 and qty <= threshold:
        return "Low stock"
    return "Available"


def load_items(ws) -> list[dict]:
    """Read master tracker rows (row 50+) into item dicts."""
    items: list[dict] = []
    today = date.today()
    for row in ws.iter_rows(min_row=MASTER_HEADER_ROW + 1, values_only=True):
        b = row[1] if len(row) > 1 else None
        if not isinstance(b, str):
            continue
        item_id = b.strip()
        if not ID_RE.match(item_id):
            continue

        desc = (str(row[0]).strip() if row[0] is not None else "")
        vendor = (str(row[2]).strip() if row[2] is not None else "")
        lead = to_int(row[3])
        inv_value = to_int(row[4])
        po_value = to_int(row[5])
        qty = to_int(row[6])
        # row[7] is "Months on Hand" — often "#DIV/0!"; treat all non-numeric as null
        months_raw = row[7]
        months: int | None
        if isinstance(months_raw, (int, float)) and months_raw == months_raw:
            months = int(months_raw)
        else:
            months = None
        po_note = (str(row[8]).strip() if row[8] is not None else "")
        po_due_raw = row[9]
        po_due = iso_date(po_due_raw) or (str(po_due_raw).strip() if po_due_raw is not None else "")
        on_order = to_int(row[10])
        must_order = to_int(row[11])
        threshold = to_int(row[12])

        prefix = item_id[:2]
        category = CATEGORY_LABEL.get(prefix, "Other")

        items.append(
            {
                "_id": item_id,
                "desc": desc,
                "id": item_id,
                "vendor": vendor,
                "lead": lead,
                "invValue": inv_value,
                "poValue": po_value,
                "qty": qty,
                "months": months,
                "poNote": po_note,
                "poDue": po_due if po_due not in ("None", "") else "",
                "onOrder": on_order,
                "mustOrder": must_order,
                "category": category,
                "threshold": threshold,
                # location / lot / expires / status are filled in after ledger load
            }
        )
    return items


def load_ledger(ws) -> dict:
    """Read a per-item sheet into a LedgerBucket-compatible dict.

    Excel layout (fixed template):
        R4   Item Number (col C), Item Name (col F)
        R5   Column headers
        R6+  Entry rows, with subtotal "Total" rows at fixed positions
             (typically rows 10, 16, 22, 26, 28).

    We collect all entry rows (skipping Totals) into a flat list, then
    flatten the bucket to the existing LedgerBucket shape (single location +
    array of distinct lots + most-recent expiration).
    """
    entries: list[dict] = []
    for row in ws.iter_rows(min_row=6, max_row=30, values_only=True):
        a = row[0]
        if a is None:
            continue
        if isinstance(a, str) and a.strip().lower() == "total":
            continue
        if not isinstance(a, (datetime, date)):
            # Some rows have non-date text; skip them
            continue
        date_iso = iso_date(a)
        units_in = to_int(row[1])
        units_out = to_int(row[2])
        balance = to_int(row[3])
        location = (str(row[4]).strip() if row[4] is not None else "")
        lot = (str(row[5]).strip() if row[5] is not None else "")
        exp = iso_date(row[6]) or ""
        initials = (str(row[7]).strip() if row[7] is not None else "")
        entries.append(
            {
                "date": date_iso,
                "in": units_in,
                "out": units_out,
                "balance": balance,
                "location": location,
                "lot": lot,
                "exp": exp,
                "initials": initials,
                "reason": "",
            }
        )

    if not entries:
        return {"location": "—", "lots": [], "exp": "—", "entries": []}

    # Sort newest first
    entries.sort(key=lambda e: e["date"] or "", reverse=True)

    # Flatten to the existing LedgerBucket shape
    latest = entries[0]
    lots: list[str] = []
    seen: set[str] = set()
    for e in entries:
        lot = e["lot"]
        if lot and lot not in seen:
            seen.add(lot)
            lots.append(lot)

    return {
        "location": latest["location"] or "—",
        "lots": lots,
        "exp": latest["exp"] or "—",
        "entries": entries,
    }


def replace_block(text: str, start_re: str, end_re: str, replacement: str) -> str:
    """Replace text between (inclusive) start_re and end_re markers."""
    m_start = re.search(start_re, text, re.MULTILINE)
    if not m_start:
        raise SystemExit(f"Marker not found: {start_re!r}")
    m_end = re.search(end_re, text[m_start.end():], re.MULTILINE)
    if not m_end:
        raise SystemExit(f"End marker not found: {end_re!r}")
    end_abs = m_start.end() + m_end.end()
    return text[: m_start.start()] + replacement + text[end_abs:]


def to_ts_literal(obj) -> str:
    """JSON.stringify-compatible output that is valid TS literal."""
    return json.dumps(obj, indent=2, ensure_ascii=False)


def main() -> None:
    ap = argparse.ArgumentParser(description="Sync seed.ts from the Demo Inventory Control workbook.")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not args.xlsx.exists():
        sys.exit(f"xlsx not found: {args.xlsx}")
    if not SEED_PATH.exists():
        sys.exit(f"seed.ts not found: {SEED_PATH}")

    print(f"[sync] reading: {args.xlsx}")
    wb = openpyxl.load_workbook(args.xlsx, data_only=True, read_only=True)

    master = wb["Component Inventory Tracker"]
    items = load_items(master)
    print(f"[sync] items: {len(items)}")

    ledgers: dict[str, dict] = {}
    item_sheet_ids = {s for s in wb.sheetnames if ID_RE.match(s)}
    today = date.today()
    matched = 0
    for it in items:
        sid = it["_id"]
        if sid in item_sheet_ids:
            ledgers[sid] = load_ledger(wb[sid])
            matched += 1
        # Backfill location / lot / expires / status from ledger (most recent)
        bucket = ledgers.get(sid)
        if bucket and bucket["entries"]:
            it["location"] = bucket["location"]
            it["lot"] = bucket["lots"][0] if bucket["lots"] else "—"
            it["expires"] = bucket["exp"]
        else:
            it["location"] = "—"
            it["lot"] = "—"
            it["expires"] = "—"
        it["status"] = derive_status(it["qty"], it["threshold"], None if it["expires"] in ("—", "") else it["expires"], today)
        del it["_id"]
    print(f"[sync] ledger sheets matched: {matched}")

    # Re-order item keys to match existing seed.ts shape exactly
    item_keys = [
        "desc", "id", "vendor", "lead", "invValue", "poValue", "qty", "months",
        "poNote", "poDue", "onOrder", "mustOrder", "category", "location", "lot",
        "expires", "threshold", "status",
    ]
    items_ordered = [{k: it[k] for k in item_keys} for it in items]

    seed_text = SEED_PATH.read_text(encoding="utf-8")

    items_block = "  const items = " + to_ts_literal(items_ordered) + ";\n"
    ledgers_block = "  const ledgers = " + to_ts_literal(ledgers) + ";\n"

    # Replace items[] block. Array close in the original is `];` at column 0.
    seed_text = replace_block(
        seed_text,
        r"^  const items = \[",
        r"^\s*\];\n",
        items_block,
    )
    # Replace ledgers{} block. Object close in the original is `};` at column 0.
    seed_text = replace_block(
        seed_text,
        r"^  const ledgers = \{",
        r"^\s*\};\n",
        ledgers_block,
    )

    if args.dry_run:
        print(f"[sync] dry-run: would write {len(seed_text):,} bytes to {SEED_PATH}")
        return

    SEED_PATH.write_text(seed_text, encoding="utf-8")
    print(f"[sync] wrote {SEED_PATH} ({len(seed_text):,} bytes)")


if __name__ == "__main__":
    main()
