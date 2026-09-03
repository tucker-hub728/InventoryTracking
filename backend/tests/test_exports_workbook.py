"""Unit tests for the tabbed Excel workbook export.

No database. Covers the things a workbook can get wrong that a CSV cannot:
Excel's worksheet-naming rules, native number/date types, and the summary tab's
status label matching the screen it was exported from.

    docker compose exec -T -e DATABASE_URL=... backend uv run pytest \\
        tests/test_exports_workbook.py -q
"""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from app.exports.xlsx_writer import (
    AWAITING_RELEASE_LABEL,
    DETAIL_HEADERS,
    SUMMARY_HEADERS,
    build_workbook,
    sheet_name_for,
    summary_status_label,
)
from openpyxl import load_workbook

GENERATED = datetime(2026, 8, 25, 12, 0, tzinfo=UTC)


# ── Worksheet naming: Excel's rules, which real SKUs can trip ───────────


def test_ordinary_sku_is_used_verbatim():
    assert sheet_name_for("CR-1042", set()) == "CR-1042"


@pytest.mark.parametrize("bad", ["CR/1042", "CR[1042]", "CR:1042", "CR*10?42", "CR\\1042"])
def test_illegal_characters_are_replaced(bad):
    """Excel refuses ``[ ] : * ? / \\`` in a sheet name — openpyxl would raise."""
    name = sheet_name_for(bad, set())
    assert not set(name) & set("[]:*?/\\")


def test_name_is_truncated_to_excel_limit():
    name = sheet_name_for("CR-" + "X" * 60, set())
    assert len(name) <= 31


def test_duplicate_names_get_a_suffix():
    """Two items must never collide — the second would overwrite the first."""
    used: set[str] = set()
    first = sheet_name_for("CR-1042", used)
    second = sheet_name_for("CR-1042", used)
    assert first != second
    assert second.startswith("CR-1042")


def test_duplicate_detection_is_case_insensitive():
    """Excel treats "cr-001" and "CR-001" as the same sheet."""
    used: set[str] = set()
    sheet_name_for("CR-001", used)
    assert sheet_name_for("cr-001", used) != "cr-001"


def test_long_names_that_share_a_prefix_stay_distinct():
    """Truncating blindly would make these the same 31 characters."""
    used: set[str] = set()
    stem = "CR-VERY-LONG-ITEM-NUMBER-PREFIX"
    first = sheet_name_for(stem + "-AAA", used)
    second = sheet_name_for(stem + "-BBB", used)
    assert first != second
    assert len(second) <= 31


def test_blank_name_falls_back():
    assert sheet_name_for("", set())


# ── Summary status label: must read like the screen ────────────────────


@pytest.mark.parametrize(
    ("code", "expected"),
    [("available", "Available"), ("low_stock", "Low stock"), ("out_of_stock", "Out of stock")],
)
def test_status_labels_match_the_ui(code, expected):
    label = summary_status_label(
        code,
        qty_available=Decimal("5"),
        reorder_threshold=Decimal("0"),
        qty_awaiting_release=0,
    )
    assert label == expected


def test_covered_shortage_reads_as_awaiting_release():
    """"Low stock" is a purchase signal; if the gap already arrived, saying it
    would send someone to buy what is in the building."""
    assert (
        summary_status_label(
            "low_stock",
            qty_available=Decimal("5"),
            reorder_threshold=Decimal("20"),
            qty_awaiting_release=15,
        )
        == AWAITING_RELEASE_LABEL
    )


def test_partly_covered_shortage_keeps_its_warning():
    """An item still short after releasing everything needs an order."""
    assert (
        summary_status_label(
            "low_stock",
            qty_available=Decimal("5"),
            reorder_threshold=Decimal("20"),
            qty_awaiting_release=3,
        )
        == "Low stock"
    )


def test_healthy_item_with_arrivals_stays_available():
    """A delivery on top of a healthy shelf is not news."""
    assert (
        summary_status_label(
            "available",
            qty_available=Decimal("500"),
            reorder_threshold=Decimal("20"),
            qty_awaiting_release=100,
        )
        == "Available"
    )


# ── The workbook itself ────────────────────────────────────────────────


def _summary_row(**over):
    base = {
        "item_number": "CR-1042",
        "item_description": "Media 500mL",
        "is_consumable": True,
        "qty_on_hand": Decimal("50"),
        "uom": "Each",
        "status": "Available",
        "qty_awaiting_release": 0,
        "qty_on_order": Decimal("0"),
        "inv_value": Decimal("47.50"),
        "po_value": Decimal("0"),
        "reorder_threshold": Decimal("20"),
        "updated_at": datetime(2026, 8, 24, 9, 30, tzinfo=UTC),
        "sheet": "CR-1042",
    }
    return base | over


def _detail_row(**over):
    base = {
        "row_kind": "on_hand",
        "lot_number": "PL-301-07",
        "lot_status": "available",
        "lot_release_state": None,
        "expiration_date": date(2027, 3, 1),
        "is_expired": False,
        "days_to_expiry": 188,
        "location_path": "Storage › Analytical Lab › Shelf A",
        "supplier_name": "Accellix",
        "po_reference": "PO-1042",
        "order_status": None,
        "qty_on_hand": 50,
        "qty_awaiting_release": None,
        "qty_on_order": None,
        "unit_price_per_consume_unit": Decimal("0.950000"),
        "unit_price_per_order_unit": None,
        "extended_value_amount": Decimal("47.500000"),
    }
    return base | over


def _build(summary_rows=None, item_tabs=None, title="Raw Materials"):
    data = build_workbook(
        summary_title=title,
        summary_rows=summary_rows if summary_rows is not None else [_summary_row()],
        item_tabs=item_tabs
        if item_tabs is not None
        else [
            {
                "sheet": "CR-1042",
                "facts": {
                    "item_number": "CR-1042",
                    "item_description": "Media 500mL",
                    "item_tier_name": "Raw Materials",
                    "uom": "Each",
                    "is_consumable": True,
                    "item_status_code": "Available",
                    "qty_on_hand": Decimal("50"),
                    "qty_available": Decimal("50"),
                    "qty_awaiting_release": 0,
                    "qty_on_order": Decimal("0"),
                    "inv_value": Decimal("47.50"),
                    "po_value": Decimal("0"),
                    "reorder_threshold": Decimal("20"),
                    "last_movement_at": datetime(2026, 8, 24, 9, 30, tzinfo=UTC),
                },
                "detail_rows": [_detail_row()],
            }
        ],
        generated_at=GENERATED,
    )
    return load_workbook(io.BytesIO(data)), data


def test_workbook_has_a_summary_tab_then_one_tab_per_item():
    book, _ = _build()
    assert book.sheetnames == ["Raw Materials", "CR-1042"]


def test_summary_tab_headers_are_the_on_screen_columns():
    book, _ = _build()
    sheet = book["Raw Materials"]
    headers = [sheet.cell(row=4, column=i + 1).value for i in range(len(SUMMARY_HEADERS))]
    assert headers == [label for label, _ in SUMMARY_HEADERS]


def test_detail_tab_headers():
    book, _ = _build()
    sheet = book["CR-1042"]
    header_row = next(
        r for r in range(1, 30) if sheet.cell(row=r, column=1).value == DETAIL_HEADERS[0][0]
    )
    headers = [sheet.cell(row=header_row, column=i + 1).value for i in range(len(DETAIL_HEADERS))]
    assert headers == [label for label, _ in DETAIL_HEADERS]


def test_numbers_are_written_as_numbers_not_text():
    """The whole reason to ship a workbook — finance can SUM without a
    text-to-columns pass first."""
    book, _ = _build()
    sheet = book["Raw Materials"]
    qty = sheet.cell(row=5, column=1 + [k for _, k in SUMMARY_HEADERS].index("qty_on_hand"))
    value = sheet.cell(row=5, column=1 + [k for _, k in SUMMARY_HEADERS].index("inv_value"))
    assert isinstance(qty.value, int | float)
    assert isinstance(value.value, float)
    assert value.value == pytest.approx(47.50)


def _price_cell(row):
    """The "Unit price (each)" cell on the first detail row of an item tab."""
    book, _ = _build(
        item_tabs=[
            {
                "sheet": "CR-1042",
                "facts": {"item_number": "CR-1042"},
                "detail_rows": [_detail_row(unit_price_per_consume_unit=row)],
            }
        ]
    )
    sheet = book["CR-1042"]
    col = 1 + [k for _, k in DETAIL_HEADERS].index("unit_price_per_consume_unit")
    header = next(
        r for r in range(1, 20) if sheet.cell(row=r, column=1).value == "Type"
    )
    return sheet.cell(row=header + 1, column=col)


@pytest.mark.parametrize(
    ("price", "rendered"),
    [
        # The complaint: an ordinary price padded out to six decimals reads as
        # false precision.
        (Decimal("4.000000"), "4.00"),
        (Decimal("0.040000"), "0.04"),
        (Decimal("0.950000"), "0.95"),
        # ...but a price that genuinely does not divide keeps every digit.
        # "$95.00 over a pack of 3" is why the column is 6dp in the first place.
        (Decimal("31.666667"), "31.666667"),
        (Decimal("0.083333"), "0.083333"),
    ],
)
def test_unit_price_shows_two_decimals_unless_it_needs_more(price, rendered):
    """Display only. Excel renders per the number format; the cell still holds
    the unrounded value, which is what a SUM or a multiplication reads."""
    cell = _price_cell(price)
    assert cell.number_format == "#,##0.00####"
    # The stored value is untouched — this is a formatting change, not a
    # rounding one.
    assert cell.value == pytest.approx(float(price))
    # What Excel will actually draw, derived from the format the same way.
    whole, _, frac = rendered.partition(".")
    assert len(frac) >= 2
    assert f"{float(price):.6f}".startswith(f"{whole}.{frac}".rstrip("0") or whole)


def test_price_format_never_hides_a_nonzero_price():
    """A format that rounded to 2dp would render 0.004 as "0.00" — a real
    price displayed as no price, next to a column where blank means unknown."""
    cell = _price_cell(Decimal("0.004000"))
    assert cell.value == pytest.approx(0.004)
    assert cell.number_format.count("#", cell.number_format.index(".")) >= 4


def test_dates_are_written_as_dates():
    book, _ = _build()
    sheet = book["Raw Materials"]
    cell = sheet.cell(row=5, column=1 + [k for _, k in SUMMARY_HEADERS].index("updated_at"))
    assert isinstance(cell.value, datetime)
    # Excel has no timezone concept — a tz-aware value would be written as text.
    assert cell.value.tzinfo is None


def test_missing_price_is_a_genuinely_empty_cell():
    """The blank-not-zero rule holds in the workbook too: an unpriced lot must
    not render as 0, which would read as 'worthless'."""
    book, _ = _build(
        item_tabs=[
            {
                "sheet": "CR-1042",
                "facts": {"item_number": "CR-1042"},
                "detail_rows": [
                    _detail_row(unit_price_per_consume_unit=None, extended_value_amount=None)
                ],
            }
        ]
    )
    sheet = book["CR-1042"]
    header_row = next(
        r for r in range(1, 30) if sheet.cell(row=r, column=1).value == DETAIL_HEADERS[0][0]
    )
    col = 1 + [k for _, k in DETAIL_HEADERS].index("extended_value_amount")
    assert sheet.cell(row=header_row + 1, column=col).value is None


def test_booleans_render_as_yes_no():
    book, _ = _build()
    sheet = book["Raw Materials"]
    cell = sheet.cell(row=5, column=1 + [k for _, k in SUMMARY_HEADERS].index("is_consumable"))
    assert cell.value == "Yes"


def test_row_kind_is_shown_in_english():
    book, _ = _build()
    sheet = book["CR-1042"]
    header_row = next(
        r for r in range(1, 30) if sheet.cell(row=r, column=1).value == DETAIL_HEADERS[0][0]
    )
    assert sheet.cell(row=header_row + 1, column=1).value == "In stock"


def test_sku_links_to_its_tab_only_when_one_exists():
    """An item with nothing to show gets a summary line but no tab, so its SKU
    must stay plain text rather than becoming a link that lands nowhere."""
    book, _ = _build(
        summary_rows=[_summary_row(), _summary_row(item_number="CR-9999", sheet=None)],
        item_tabs=[],
    )
    sheet = book["Raw Materials"]
    assert sheet.cell(row=5, column=1).hyperlink is not None
    assert sheet.cell(row=6, column=1).hyperlink is None


def test_sku_link_is_internal_and_points_at_the_right_tab():
    """The link must carry ``location``, not ``target``.

    Assigning a plain string to ``cell.hyperlink`` sets the attribute and looks
    correct, but openpyxl turns it into an *external* relationship — Excel
    treats it as a URL outside the file and clicking does nothing. Asserting
    "a hyperlink exists" is not enough to catch that; this asserts the shape.
    """
    book, _ = _build()
    link = book["Raw Materials"].cell(row=5, column=1).hyperlink
    assert link.location == "'CR-1042'!A1"
    assert link.target is None


def test_back_link_on_an_item_tab_is_internal():
    book, _ = _build()
    link = book["CR-1042"]["A2"].hyperlink
    assert link.location == "'Raw Materials'!A1"
    assert link.target is None


def test_workbook_contains_no_external_relationships():
    """The regression guard, at the file level.

    An external hyperlink relationship is exactly what a string assignment
    produces, and it is invisible from the openpyxl object model once the file
    round-trips — so check the XML the recipient actually opens.
    """
    _, data = _build()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.endswith(".rels"):
                assert b'TargetMode="External"' not in archive.read(name), name


def test_internal_link_survives_a_sheet_name_with_a_space():
    """The summary tab is normally two words; an unquoted name breaks the ref."""
    book, _ = _build()
    assert book["CR-1042"]["A2"].hyperlink.location == "'Raw Materials'!A1"


def test_apostrophe_in_a_sheet_name_is_doubled():
    """Excel escapes an apostrophe inside a quoted sheet name by doubling it;
    a single one would truncate the reference."""
    book, _ = _build(
        summary_rows=[_summary_row(sheet="O'Brien")],
        item_tabs=[
            {"sheet": "O'Brien", "facts": {"item_number": "CR-1042"}, "detail_rows": []}
        ],
    )
    assert book["Raw Materials"].cell(row=5, column=1).hyperlink.location == "'O''Brien'!A1"


def test_formula_injection_is_neutralised_in_the_workbook():
    """Excel executes a leading '=' loaded from .xlsx exactly as from .csv."""
    book, _ = _build(summary_rows=[_summary_row(item_description="=1+1")])
    sheet = book["Raw Materials"]
    col = 1 + [k for _, k in SUMMARY_HEADERS].index("item_description")
    assert sheet.cell(row=5, column=col).value.startswith("'")


def test_empty_export_still_produces_a_valid_workbook():
    book, data = _build(summary_rows=[], item_tabs=[])
    assert book.sheetnames == ["Raw Materials"]
    assert data[:2] == b"PK"  # a real zip container, i.e. a loadable .xlsx


def test_illegal_characters_in_the_title_are_cleaned():
    book, _ = _build(title="Raw/Materials")
    assert not set(book.sheetnames[0]) & set("[]:*?/\\")
